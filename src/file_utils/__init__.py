from __future__ import annotations

from pathlib import Path
from typing import Callable, Dict, Union
import csv
import logging
import mimetypes
import tempfile
from PIL import Image, ImageOps


# Опциональные зависимости
try:
    import fitz  # PyMuPDF
except ImportError:  # pragma: no cover - dependency missing
    fitz = None

try:
    from docx import Document
except ImportError:  # pragma: no cover - dependency missing
    Document = None

try:
    import xlrd
except ImportError:  # pragma: no cover - dependency missing
    xlrd = None

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency missing
    openpyxl = None

try:
    import magic
except ImportError:  # pragma: no cover - optional dependency
    magic = None

# OCR для изображений — модуль может отсутствовать
try:
    from .image_ocr import extract_text_image  # ожидается сигнатура (path: Path, language: str) -> str
except Exception:  # pragma: no cover - optional module
    extract_text_image = None  # type: ignore

from .mrz import parse_mrz

logger = logging.getLogger(__name__)


class UnsupportedFileType(ValueError):
    """Выбрасывается, когда тип файла не поддерживается или не определён."""


_PARSER_REGISTRY: Dict[str, Callable[[Path], str]] = {}


def register_parser(ext: str) -> Callable[[Callable[[Path], str]], Callable[[Path], str]]:
    """Декоратор для регистрации парсеров файлов."""

    def decorator(func: Callable[[Path], str]) -> Callable[[Path], str]:
        _PARSER_REGISTRY[ext.lower()] = func
        return func

    return decorator


# ---------- Парсеры для «текстовых» форматов ----------


@register_parser(".txt")
def extract_text_txt(path: Path) -> str:
    """Извлечение текста из .txt файла."""
    return path.read_text(encoding="utf-8")


@register_parser(".md")
def extract_text_md(path: Path) -> str:
    """Извлечение текста из .md файла."""
    return path.read_text(encoding="utf-8")


@register_parser(".pdf")
def extract_text_pdf(path: Path, language: str = "eng") -> str:
    """Извлечение текста из PDF с помощью PyMuPDF.

    Если страница не содержит текстового слоя, при наличии ``extract_text_image``
    страница будет конвертирована в изображение и обработана через OCR.

    :param path: путь к PDF-файлу.
    :param language: язык OCR для страниц без текстового слоя.
    :return: извлечённый текст по всем страницам.
    :raises RuntimeError: если PyMuPDF или OCR-модуль недоступны.
    """
    if fitz is None:
        raise RuntimeError("PyMuPDF не установлен, обработка PDF невозможна")

    parts: list[str] = []
    with fitz.open(path) as doc:
        for page in doc:
            text = page.get_text()
            if text.strip():
                parts.append(text)
                continue

            if extract_text_image is None:
                raise RuntimeError(
                    "OCR недоступен: функция extract_text_image не найдена"
                )

            pix = page.get_pixmap()

            # На Windows невозможно перезаписать открытый временный файл,
            # поэтому сначала закрываем его, а затем сохраняем изображение.
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
            tmp_path = Path(tmp.name)
            tmp.close()

            try:
                pix.save(tmp_path)
                ocr_text = extract_text_image(tmp_path, language=language)
            finally:
                tmp_path.unlink(missing_ok=True)
            if ocr_text.strip():
                parts.append(ocr_text)

    return "\n".join(parts)


@register_parser(".docx")
def extract_text_docx(path: Path) -> str:
    """Извлечение текста из DOCX (python-docx)."""
    if Document is None:
        raise RuntimeError("python-docx не установлен")
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


@register_parser(".csv")
def extract_text_csv(path: Path) -> str:
    """Извлечение текста из CSV (через csv.reader)."""
    lines: list[str] = []
    with path.open(newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row in reader:
            lines.append(",".join(row))
    return "\n".join(lines)


@register_parser(".xls")
def extract_text_xls(path: Path) -> str:
    """Извлечение текста из XLS (через xlrd)."""
    if xlrd is None:
        raise RuntimeError("xlrd не установлен")
    book = xlrd.open_workbook(path)
    lines: list[str] = []
    for sheet in book.sheets():
        for row_idx in range(sheet.nrows):
            row = sheet.row_values(row_idx)
            lines.append(",".join(str(cell) for cell in row))
    return "\n".join(lines)


@register_parser(".xlsx")
def extract_text_xlsx(path: Path) -> str:
    """Извлечение текста из XLSX (через openpyxl)."""
    if openpyxl is None:
        raise RuntimeError("openpyxl не установлен")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row]
            lines.append(",".join(values))
    wb.close()
    return "\n".join(lines)


# ---------- Входная точка с поддержкой OCR ----------

def extract_text(file_path: Union[str, Path], language: str = "eng") -> str:
    """
    Извлечь текст из поддерживаемого файла.

    Поддерживаемые форматы:
      - Текстовые: .txt, .md, .pdf, .docx
      - Таблицы: .csv, .xls, .xlsx
      - Изображения (OCR): .jpg, .jpeg, .png, .tiff (через image_ocr.extract_text_image)

    :param file_path: путь к файлу.
    :param language: язык OCR (ISO-коды tesseract, напр. 'eng', 'rus', 'deu').
    :return: извлечённый текст.
    :raises UnsupportedFileType: если расширение не поддерживается или не определено.
    :raises RuntimeError: если требуемая зависимость для формата не установлена.
    """
    path = Path(file_path)
    ext = path.suffix.lower()
    logger.info("Extracting text from %s", path)

    if not ext:
        mime = None
        if magic is not None:
            try:
                mime = magic.from_file(str(path), mime=True)
            except Exception:  # pragma: no cover - magic failure
                mime = None
        if mime:
            ext = mimetypes.guess_extension(mime) or ""
            logger.debug("Guessed extension %s for %s", ext, path)
        if not ext:
            logger.error("Cannot determine file type for %s", path)
            raise UnsupportedFileType("Не удалось определить тип файла")

    # Ветвь для изображений — нужен отдельный параметр language
    if ext in {".jpg", ".jpeg", ".png", ".tiff"}:
        if extract_text_image is None:
            logger.error("OCR module unavailable for %s", path)
            raise RuntimeError("Модуль OCR недоступен: .image_ocr.extract_text_image не найден")
        text = extract_text_image(path, language=language)
        logger.debug("Extracted %d characters from %s via OCR", len(text), path)
        return text

    if ext == ".pdf":
        text = extract_text_pdf(path, language=language)
        logger.debug("Extracted %d characters from %s", len(text), path)
        return text

    # Обычные «текстовые» форматы
    parser = _PARSER_REGISTRY.get(ext)
    if parser is None:
        logger.error("Unsupported/unknown file extension: %s", ext)
        raise UnsupportedFileType(f"Unsupported/unknown file extension: {ext}")
    text = parser(path)
    logger.debug("Extracted %d characters from %s", len(text), path)
    return text


# ---------- Вспомогательные утилиты ----------

def merge_images_to_pdf(paths: list[Path]) -> Path:
    """Преобразовать несколько изображений в один PDF во временном файле.

    Обрабатываются изображения в различных цветовых пространствах и размерах.

    :param paths: список путей к изображениям.
    :return: путь к созданному временному PDF-файлу.
    """
    if not paths:
        raise ValueError("No images provided")

    images: list[Image.Image] = []
    max_w = max_h = 0
    for path in paths:
        with Image.open(path) as img:
            if img.mode != "RGB":
                img = img.convert("RGB")
            img.load()
            max_w = max(max_w, img.width)
            max_h = max(max_h, img.height)
            images.append(img)

    normalized = [
        ImageOps.pad(img, (max_w, max_h), color=(255, 255, 255))
        if img.size != (max_w, max_h)
        else img
        for img in images
    ]

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    first, *rest = normalized
    first.save(tmp.name, save_all=True, append_images=rest, format="PDF")
    tmp_path = Path(tmp.name)
    tmp.close()
    return tmp_path


__all__ = [
    "UnsupportedFileType",
    "extract_text",
    "register_parser",
    "extract_text_txt",
    "extract_text_md",
    "extract_text_pdf",
    "extract_text_docx",
    "extract_text_csv",
    "extract_text_xls",
    "extract_text_xlsx",
    "merge_images_to_pdf",
    "parse_mrz",
    "translate_text",
    "load_plugins",
]

def load_plugins() -> None:
    """Автоматически обнаружить и загрузить плагины, если они доступны."""
    try:
        import importlib
        plugin_module = importlib.import_module("plugins")
    except Exception:  # pragma: no cover - отсутствие плагинов не критично
        logger.debug("Plugin module not found", exc_info=True)
        return
    try:
        plugin_module.load_plugins()  # type: ignore[attr-defined]
    except Exception:  # pragma: no cover - ошибки плагинов не критичны
        logger.warning("Plugin loading failed", exc_info=True)


async def translate_text(
    text: str,
    target_lang: str,
    *,
    model: str | None = None,
    api_key: str | None = None,
    base_url: str | None = None,
) -> str:
    """Перевести *text* на язык ``target_lang`` с помощью OpenRouter."""

    from services.openrouter import OpenRouterError, chat

    prompt = f"Translate the following text to {target_lang}:\n{text}"
    try:
        reply, _, _ = await chat(
            messages=[{"role": "user", "content": prompt}],
            model=model,
            api_key=api_key,
            base_url=base_url,
        )
    except OpenRouterError as exc:
        logger.error("Translation request failed: %s", exc)
        raise RuntimeError("Translation request failed") from exc
    return reply
