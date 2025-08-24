import sys
from pathlib import Path
import csv
import importlib.util
from tempfile import TemporaryDirectory
import unittest
import shutil
from PIL import Image, ImageDraw, ImageFont

sys.path.append(str(Path(__file__).resolve().parents[1] / "src"))

from file_utils import extract_text, merge_images_to_pdf


def has_module(name: str) -> bool:
    return importlib.util.find_spec(name) is not None


def _create_file(base: Path, ext: str, text: str) -> Path:
    path = base / f"file{ext}"
    if ext in {".txt", ".md"}:
        path.write_text(text, encoding="utf-8")
    elif ext == ".csv":
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(text.split(","))
    elif ext == ".pdf":
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), text)
        doc.save(path)
        doc.close()
    elif ext == ".docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph(text)
        doc.save(path)
    elif ext == ".xlsx":
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(text.split(","))
        wb.save(path)
    elif ext == ".xls":
        import xlwt

        wb = xlwt.Workbook()
        ws = wb.add_sheet("Sheet1")
        for col, val in enumerate(text.split(",")):
            ws.write(0, col, val)
        wb.save(path)
    else:
        raise ValueError("unsupported ext")
    return path


class TestExtractText(unittest.TestCase):
    def test_txt(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".txt", "hello txt")
            self.assertEqual(extract_text(path).strip(), "hello txt")

    def test_md(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".md", "hello md")
            self.assertEqual(extract_text(path).strip(), "hello md")

    def test_csv(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".csv", "a,b")
            self.assertEqual(extract_text(path).strip(), "a,b")

    @unittest.skipUnless(has_module("openpyxl"), "openpyxl not installed")
    def test_xlsx(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".xlsx", "a,b")
            self.assertEqual(extract_text(path).strip(), "a,b")

    @unittest.skipUnless(has_module("xlrd") and has_module("xlwt"), "xlrd/xlwt not installed")
    def test_xls(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".xls", "a,b")
            self.assertEqual(extract_text(path).strip(), "a,b")

    @unittest.skipUnless(has_module("fitz"), "PyMuPDF not installed")
    def test_pdf(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".pdf", "hello pdf")
            self.assertEqual(extract_text(path).strip(), "hello pdf")

    @unittest.skipUnless(has_module("docx"), "python-docx not installed")
    def test_docx(self):
        with TemporaryDirectory() as tmp:
            path = _create_file(Path(tmp), ".docx", "hello docx")
            self.assertEqual(extract_text(path).strip(), "hello docx")

    @unittest.skipUnless(shutil.which("tesseract"), "tesseract not installed")
    def test_png(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.png"
            image = Image.new("RGB", (100, 50), color="white")
            draw = ImageDraw.Draw(image)
            draw.text((10, 10), "PNG", fill="black", font=ImageFont.load_default())
            image.save(path)
            self.assertIn("PNG", extract_text(path))

    @unittest.skipUnless(shutil.which("tesseract"), "tesseract not installed")
    def test_tiff(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.tiff"
            image = Image.new("RGB", (100, 50), color="white")
            draw = ImageDraw.Draw(image)
            draw.text((10, 10), "TIFF", fill="black", font=ImageFont.load_default())
            image.save(path)
            self.assertIn("TIFF", extract_text(path))


@unittest.skipUnless(has_module("fitz"), "PyMuPDF not installed")
class TestMergeImagesToPdf(unittest.TestCase):
    def test_merge_images_to_pdf(self):
        with TemporaryDirectory() as tmp:
            base = Path(tmp)
            img1 = base / "img1.jpg"
            Image.new("RGB", (100, 100), color="red").save(img1)
            img2 = base / "img2.png"
            Image.new("RGBA", (80, 120), color=(0, 255, 0, 128)).save(img2)
            pdf = merge_images_to_pdf([img1, img2])
            self.assertTrue(pdf.exists())
            import fitz
            with fitz.open(pdf) as doc:
                self.assertEqual(doc.page_count, 2)
            pdf.unlink()


if __name__ == "__main__":
    unittest.main()
